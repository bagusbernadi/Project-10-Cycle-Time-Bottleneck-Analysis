"""06 - Cek silang SQL vs Excel (Q14). Workbook memakai SUMIFS, dihitung ulang oleh LibreOffice.

Output: project/results/2-analysis/q14_excel_crosscheck.xlsx (sudah dihitung ulang)
        project/results/2-analysis/05_crosscheck_spreadsheet.csv
"""
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from common import CLEAN_DIR, RES_ANALYSIS  # noqa: E402

o = pd.read_csv(CLEAN_DIR / "daily_output_clean.csv")
sql = pd.read_csv(RES_ANALYSIS / "q14_agg.csv")
rng = pd.Series(range(len(sql)))
pick = pd.concat([
    sql[sql["operator_id"].isin(["Op_03", "Op_22"])].head(4),
    sql.sample(10, random_state=20260302),
]).drop_duplicates(["station", "operator_id", "model"]).reset_index(drop=True)

wb = openpyxl.Workbook()
data = wb.active
data.title = "data"
cols = ["date", "station", "operator_id", "model", "run_hours", "ok_qty"]
data.append(cols)
for r in o[cols].itertuples(index=False):
    data.append(list(r))
n = len(o) + 1
for c in data[1]:
    c.font = Font(bold=True)

ws = wb.create_sheet("cek")
head = ["station", "operator_id", "model", "total_jam (SUMIFS)", "total_ok (SUMIFS)",
        "CT Excel (s)", "CT SQL (s)", "selisih", "status"]
ws.append(head)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="1A73E8")
for i, r in enumerate(pick.itertuples(index=False), start=2):
    ws.append([r.station, r.operator_id, r.model,
               f'=SUMIFS(data!$E$2:$E${n},data!$B$2:$B${n},A{i},data!$C$2:$C${n},B{i},data!$D$2:$D${n},C{i})',
               f'=SUMIFS(data!$F$2:$F${n},data!$B$2:$B${n},A{i},data!$C$2:$C${n},B{i},data!$D$2:$D${n},C{i})',
               f"=ROUND(D{i}*3600/E{i},2)", round(float(r.ct_agregat_s), 2),
               f"=ABS(F{i}-G{i})", f'=IF(H{i}<0.01,"PASS","FAIL")'])
ws.column_dimensions["A"].width = 34
for col in "BCDEFGHI":
    ws.column_dimensions[col].width = 18

tmp = Path(tempfile.mkdtemp(prefix="xcheck_"))
src = tmp / "q14_excel_crosscheck.xlsx"
wb.save(src)
subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(tmp / "out"), str(src)],
               check=True, capture_output=True, timeout=300)
recalc = tmp / "out" / src.name
wb2 = openpyxl.load_workbook(recalc, data_only=True)
vals = [[c.value for c in row] for row in wb2["cek"].iter_rows(min_row=2)]
res = pd.DataFrame(vals, columns=head)
res = res[["station", "operator_id", "model", "CT Excel (s)", "CT SQL (s)", "selisih", "status"]]
res = res.rename(columns={"CT Excel (s)": "ct_excel_s", "CT SQL (s)": "ct_sql_s", "selisih": "selisih_s"})
res.to_csv(RES_ANALYSIS / "05_crosscheck_spreadsheet.csv", index=False)
shutil.copy(recalc, RES_ANALYSIS / "q14_excel_crosscheck.xlsx")
shutil.rmtree(tmp, ignore_errors=True)
print(res["status"].value_counts().to_dict(), "baris:", len(res), "selisih maks:", res["selisih_s"].max())
